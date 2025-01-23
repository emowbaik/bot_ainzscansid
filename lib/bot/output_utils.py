import os
import openpyxl
import logging
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

def generate_output_lapor(filename: str, sheet_title: str, data: list):
    """
    Membuat file Excel dari data yang diberikan.
    :param filename: Nama file output.
    :param sheet_title: Judul sheet Excel.
    :param data: Data yang akan ditulis ke dalam Excel.
    :return: Path file yang dibuat.
    """
    try:
        # Buat workbook
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = sheet_title

        # Header kolom
        headers = ["ID", "Judul Komik", "Chapter", "Posisi", "Username", "Reported At"]
        sheet.append(headers)

        for row in data:
            row_data = [
                row['id'],
                row['judul_name'],
                row['chapter'],
                row['posisi_name'],
                row['reporter_name'],
                row['reported_at'].strftime('%Y-%m-%d %H:%M:%S')  # Format the datetime to string if needed
            ]
            sheet.append(row_data)
            
        # Pastikan jumlah data sesuai header
        if any(len(row) != len(headers) for row in data):
            raise ValueError("Jumlah elemen dalam data tidak sesuai dengan header.")

        # Terapkan format tabel
        table_ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"  # Area tabel
        table = Table(displayName="ReportTable", ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium9",  # Gaya tabel
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True,
        )
        table.tableStyleInfo = style
        sheet.add_table(table)

        # Terapkan format header
        header_row = sheet[1]
        for cell in header_row:
            cell.font = Font(bold=True)  # Buat teks header bold
            cell.alignment = Alignment(horizontal="center", vertical="center")  # Atur header ke tengah

        # Terapkan border ke seluruh tabel
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        for row in sheet.iter_rows(min_row=1, max_row=len(data) + 1, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border

        # Atur lebar kolom secara otomatis berdasarkan konten terpanjang
        for col_num, column_cells in enumerate(sheet.columns, start=1):
            max_length = 0
            column_letter = get_column_letter(col_num)
            for cell in column_cells:
                try:
                    # Hitung panjang maksimum dari konten sel (header + data)
                    max_length = max(max_length, len(str(cell.value)) if cell.value else 0)
                except Exception as e:
                    print(f"Error calculating column width: {e}")
            adjusted_width = max_length + 5  # Tambahkan padding
            sheet.column_dimensions[column_letter].width = adjusted_width

        # Simpan file
        os.makedirs("output", exist_ok=True)
        filepath = os.path.join("output", filename)
        workbook.save(filepath)
        logging.debug(f"Nama kolom dalam data: {list(data[0].keys()) if data else 'Data kosong'}")

        return filepath
    except Exception as e:
        print(f"Error generating Excel report: {e}")
        return None

def generate_output_rate(filepath: str, sheet_title: str, bulan: str, user_report_count: dict, total_pendapatan: dict):
    """
    Membuat file Excel untuk laporan output rate dan menyimpannya ke lokasi yang diberikan.
    :param filepath: Lokasi file untuk menyimpan laporan.
    :param sheet_title: Judul sheet di dalam file Excel.
    :param bulan: Nama bulan laporan.
    :param user_report_count: Dictionary dengan jumlah laporan per user.
    :param total_pendapatan: Dictionary dengan total pendapatan per user.
    """
    try:
        # Buat workbook
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = sheet_title

        # Header
        headers = ["Nama User", "Jumlah Laporan", "Total Pendapatan"]
        sheet.append(headers)

        # Data
        for user, jumlah in user_report_count.items():
            sheet.append([user, jumlah, f"Rp {total_pendapatan[user]}000,00"])

        # Terapkan format header
        header_row = sheet[1]
        for cell in header_row:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Terapkan border ke seluruh tabel
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        for row in sheet.iter_rows(min_row=1, max_row=len(user_report_count) + 1, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border

        # Atur lebar kolom secara otomatis berdasarkan konten terpanjang
        for col_num, column_cells in enumerate(sheet.columns, start=1):
            max_length = max(len(str(cell.value)) for cell in column_cells if cell.value) + 5
            sheet.column_dimensions[get_column_letter(col_num)].width = max_length

        # Terapkan gaya tabel (TableStyleMedium9)
        table_ref = f"A1:{get_column_letter(len(headers))}{len(user_report_count) + 1}"
        table = Table(displayName="ReportTable", ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True,
        )
        table.tableStyleInfo = style
        sheet.add_table(table)

        # Simpan ke lokasi yang diberikan
        workbook.save(filepath)

    except Exception as e:
        raise RuntimeError(f"Terjadi kesalahan saat membuat laporan Excel: {e}")
